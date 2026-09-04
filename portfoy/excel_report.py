"""Pandas + openpyxl ile bicimlendirilmis Excel raporu."""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config
from .analytics import PortfolioAnalysis, to_dataframe
from .formatting import fmt_points

logger = logging.getLogger(__name__)

SHEET_NAME = "Portföy"

# Bicim tanimlari
_FMT_UNITS = "#,##0.######"
_FMT_PRICE = "#,##0.000000"
_FMT_MONEY = '#,##0.00 "₺"'
_FMT_PCT = '0.00"%"'
_FMT_POINTS = '+0.00;-0.00;0.00'
# Isaretli para: kazanc/kayip sutunlarinda yon tek bakista okunsun.
_FMT_MONEY_CHANGE = '+#,##0.00 "₺";-#,##0.00 "₺";0.00 "₺"'

COLUMN_FORMATS = {
    "Adet": (_FMT_UNITS, 14),
    "Güncel Fiyat": (_FMT_PRICE, 15),
    "Toplam Değer": (_FMT_MONEY, 18),
    "Toplam Maliyet": (_FMT_MONEY, 18),
    "Kar/Zarar (₺)": (_FMT_MONEY_CHANGE, 18),
    "Kar/Zarar (%)": (_FMT_PCT, 15),
    "Portföy Ağırlığı (%)": (_FMT_PCT, 19),
    "Günlük Getiri (%)": (_FMT_PCT, 17),
    "Günlük Değişim (₺)": (_FMT_MONEY_CHANGE, 18),
    "Haftalık Getiri (%)": (_FMT_PCT, 18),
    "Aylık Getiri (%)": (_FMT_PCT, 17),
    "Aylık Katkı (yp)": (_FMT_POINTS, 16),
    "Tarih": (None, 12),
    "Alış Tarihi": (None, 13),
    "Fon Kodu": (None, 11),
    "Fon Adı": (None, 46),
    "Not": (None, 48),
}

# Isaret rengi (yesil/kirmizi) uygulanacak sutunlar bu parcalarla eslesir.
_SIGNED_COLUMNS = ("Getiri", "Katkı", "Değişim", "Kar/Zarar")

_HEADER_FILL = PatternFill("solid", fgColor="FF1A1A19")
_SUMMARY_FILL = PatternFill("solid", fgColor="FFF0EFEC")
_HEADER_FONT = Font(color="FFFFFFFF", bold=True, size=11)
_HAIRLINE = Side(style="thin", color="FFE1E0D9")
_BORDER = Border(bottom=_HAIRLINE)
_GREEN = Font(color="FF006300")
_RED = Font(color="FFD03B3B")


def build_report(
    analysis: PortfolioAnalysis,
    output_dir: Path | None = None,
    filename: str | None = None,
) -> Path | None:
    """Analizi .xlsx dosyasina yazar. Yazilan dosyanin yolunu dondurur."""
    if not analysis.rows:
        logger.warning("Excel raporu icin satir yok; atlandi.")
        return None

    output_dir = output_dir or config.OUTPUT_DIR
    output_dir.mkdir(parents=True, exist_ok=True)
    filename = filename or f"portfoy_raporu_{analysis.as_of:%Y%m%d}.xlsx"
    path = output_dir / filename

    frame = to_dataframe(analysis)

    try:
        with pd.ExcelWriter(path, engine="openpyxl") as writer:
            frame.to_excel(writer, sheet_name=SHEET_NAME, index=False, startrow=0)
            sheet = writer.sheets[SHEET_NAME]
            _style_table(sheet, frame)
            next_row = len(frame) + 3
            next_row = _write_summary(sheet, frame, analysis, next_row)
            _write_warnings(sheet, analysis, next_row + 1)
    except (OSError, PermissionError) as exc:
        logger.error("Excel dosyasi yazilamadi (%s): %s", path, exc)
        return None

    return path


# --------------------------------------------------------------------------
def _style_table(sheet, frame: pd.DataFrame) -> None:
    headers = list(frame.columns)

    for index, name in enumerate(headers, start=1):
        letter = get_column_letter(index)
        cell = sheet.cell(row=1, column=index)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        number_format, width = COLUMN_FORMATS.get(name, (None, 16))
        sheet.column_dimensions[letter].width = width

        for row_index in range(2, len(frame) + 2):
            data_cell = sheet.cell(row=row_index, column=index)
            data_cell.border = _BORDER
            if number_format:
                data_cell.number_format = number_format
            if name in ("Fon Kodu", "Tarih", "Alış Tarihi"):
                data_cell.alignment = Alignment(horizontal="center")
            if name == "Not":
                data_cell.font = Font(italic=True, color="FF898781")
            # Getiri / degisim sutunlarinda isaret rengi
            if any(token in name for token in _SIGNED_COLUMNS):
                value = data_cell.value
                if isinstance(value, (int, float)):
                    data_cell.font = _GREEN if value >= 0 else _RED

    sheet.row_dimensions[1].height = 30
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(frame) + 1}"


def _write_summary(sheet, frame: pd.DataFrame, analysis: PortfolioAnalysis, row: int) -> int:
    headers = list(frame.columns)
    label_col = 1
    # Ozet degerleri "Toplam Değer" sutunuyla ayni hizada dursun.
    value_col = headers.index("Toplam Değer") + 1 if "Toplam Değer" in headers else 6

    gap = analysis.target_gap
    if gap is None:
        status = "Aylık veri eksik — hedef değerlendirilemedi"
    elif gap >= 0:
        status = f"HEDEF TUTTU ({fmt_points(gap)} yüzde puan)"
    else:
        status = f"HEDEFİN ALTINDA ({fmt_points(gap)} yüzde puan)"

    entries: list[tuple[str, object, str | None]] = [
        ("ÖZET", None, None),
        ("Toplam Portföy Değeri", analysis.total_value, _FMT_MONEY),
        ("Önceki Güne Göre Değişim", analysis.value_change("gunluk"), _FMT_MONEY_CHANGE),
        (
            "Ağırlıklı Ortalama Günlük Getiri",
            analysis.weighted_returns.get("gunluk"),
            _FMT_PCT,
        ),
        (
            "Ağırlıklı Ortalama Haftalık Getiri",
            analysis.weighted_returns.get("haftalik"),
            _FMT_PCT,
        ),
        (
            "Ağırlıklı Ortalama Aylık Getiri",
            analysis.weighted_returns.get("aylik"),
            _FMT_PCT,
        ),
        ("Toplam Maliyet", analysis.total_cost, _FMT_MONEY),
        ("Toplam Kar/Zarar", analysis.total_profit, _FMT_MONEY_CHANGE),
        ("Hedef Aylık Getiri", analysis.target_monthly_return, _FMT_PCT),
        ("Hedefe Fark (yüzde puan)", gap, _FMT_POINTS),
        ("Haftalık Değer Değişimi", analysis.value_change("haftalik"), _FMT_MONEY_CHANGE),
        ("Aylık Değer Değişimi", analysis.value_change("aylik"), _FMT_MONEY_CHANGE),
        # Akis, getiriden ayri bir satir: portfoyden cikan/giren para kar/zarar
        # degildir ama toplam degerin neden degistigini aciklayan tek sey odur.
        ("Günlük Dış Para Akışı", analysis.flow_note("gunluk") or "—", None),
        ("Aylık Dış Para Akışı", analysis.flow_note("aylik") or "—", None),
        ("Kapanan Pozisyon", analysis.closed_note or "—", None),
        # Kolonlarin tabani: "%" ile "₺" ayni seyi olcmuyor ve Excel'de sayisal
        # hucreye isaret konamadigi icin aciklamanin yeri burasi.
        ("Kolonlar", config.COLUMN_BASIS_NOTE, None),
        ("İşaretler", config.FLOW_MARK_NOTE, None),
        ("Toplam Getirinin Tabanı", config.TOTAL_BASIS_NOTE, None),
        ("Durum", status, None),
        ("Rapor Tarihi", datetime.now().strftime("%d.%m.%Y %H:%M"), None),
    ]

    for offset, (label, value, number_format) in enumerate(entries):
        current = row + offset
        label_cell = sheet.cell(row=current, column=label_col, value=label)
        label_cell.font = Font(bold=True, size=12 if label == "ÖZET" else 11)
        label_cell.fill = _SUMMARY_FILL

        value_cell = sheet.cell(row=current, column=value_col)
        value_cell.fill = _SUMMARY_FILL
        if value is None and label != "ÖZET":
            value_cell.value = "veri yok"
        elif value is not None:
            value_cell.value = value
            if number_format:
                value_cell.number_format = number_format
        value_cell.alignment = Alignment(horizontal="right")
        value_cell.font = Font(bold=True)

        # Isaretli degerler (fark ve TL degisimleri) yesil/kirmizi
        if isinstance(value, (int, float)) and (
            label.startswith("Hedefe Fark") or "Değişim" in label or "Kar/Zarar" in label
        ):
            value_cell.font = Font(bold=True, color="FF006300" if value >= 0 else "FFD03B3B")
        if label == "Durum":
            value_cell.alignment = Alignment(horizontal="left")
            sheet.cell(row=current, column=value_col).font = Font(
                bold=True,
                color="FF006300" if analysis.target_met else "FFD03B3B",
            )

        # Ozet blogunu bosluklarla birlikte gri seride tut
        for column in range(label_col, value_col + 1):
            sheet.cell(row=current, column=column).fill = _SUMMARY_FILL

    return row + len(entries)


def _write_warnings(sheet, analysis: PortfolioAnalysis, row: int) -> None:
    row = _write_partial_notes(sheet, analysis, row)

    if not analysis.failures:
        return
    title = sheet.cell(row=row, column=1, value="VERİSİ ALINAMAYAN FONLAR")
    title.font = Font(bold=True, color="FFD03B3B")
    for offset, (code, message) in enumerate(sorted(analysis.failures.items()), start=1):
        sheet.cell(row=row + offset, column=1, value=code).font = Font(bold=True)
        sheet.cell(row=row + offset, column=2, value=message)


def _write_partial_notes(sheet, analysis: PortfolioAnalysis, row: int) -> int:
    """Periyodun tamaminda portfoyde olmayan fonlari acikca yazar.

    Bu fonlarin getirisi alis tarihinden itibaren hesaplandi; okuyucu rakami
    TEFAS'in tam periyot getirisiyle karsilastirdiginda farki bilsin.
    """
    partial = analysis.partial_rows()
    if not partial:
        return row

    title = sheet.cell(row=row, column=1, value="KISMİ PERİYOT (alıştan itibaren)")
    title.font = Font(bold=True, color="FFEDA100")
    for offset, fund in enumerate(partial, start=1):
        sheet.cell(row=row + offset, column=1, value=fund.code).font = Font(bold=True)
        sheet.cell(row=row + offset, column=2, value=fund.partial_note)
    return row + len(partial) + 2
